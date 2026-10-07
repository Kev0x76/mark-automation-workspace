from __future__ import annotations

import hashlib
import json
import posixpath
import re
import shutil
import unicodedata
import xml.etree.ElementTree as ET
from collections import Counter
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .contract_snapshot import build_contract_snapshot, snapshot_item
from .contract_item_region import identify_contract_item_region, inspect_contract_row
from .delivery import normalize_original_contract_filename
from .ean13 import is_valid_ean13
from .lineart_groups import build_lineart_groups_report
from .schema_validation import write_standard_json
from ..settings import workspace_root_from_app


DISPIMG_RE = re.compile(r'DISPIMG\("([^"]+)"', re.IGNORECASE)
SUPPORTED_CONTRACT_EXTENSIONS = {".xlsx", ".xls"}
CONTRACT_MEDIA_TYPES = {
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".xls": "application/vnd.ms-excel",
}
MAX_XLSX_ARCHIVE_FILES = 20_000
MAX_XLSX_UNCOMPRESSED_BYTES = 1024 * 1024 * 1024
MAX_XLSX_ENTRY_BYTES = 256 * 1024 * 1024
MAX_XLSX_COMPRESSION_RATIO = 1_000
NAMESPACES = {
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}

FIELD_ALIASES = {
    "item_no": [
        "货号", "电脑编号", "电脑编码", "商品编号", "产品编号", "款号",
        "item no", "item number", "sku", "sku no", "article no",
    ],
    "barcode": [
        "条形码", "条码", "ean", "ean13", "ean-13", "gtin", "barcode",
        "外箱条码", "外箱条形码",
    ],
    "cn_name": [
        "中文品名", "品名", "中文", "中文名称", "产品名称", "商品名称",
        "cn product name", "chinese name", "cn name",
    ],
    "it_name": [
        "意大利文品名", "意文品名", "意语品名", "意大利语", "意大利", "意文",
        "italian name", "italiano", "italian", "it name",
    ],
    "es_name": [
        "西班牙文品名", "西文品名", "西语品名", "西班牙语", "西班牙", "西文",
        "spanish name", "espanol", "español", "spanish", "es name",
    ],
    "en_name": [
        "英文品名", "英语品名", "英文名称", "英文", "英语",
        "english name", "english", "en name",
    ],
    "qty": [
        "装箱数", "装箱数（卡）", "装箱数卡", "装箱数量", "新装箱数",
        "q.ty", "qty", "quantity", "packing qty", "qty/ctn", "pcs/ctn",
    ],
}
IMAGE_ALIASES = ["产品图", "产品图片", "图片", "图", "product image", "productimage", "image"]
REQUIRED_FIELDS = ["item_no", "barcode", "cn_name", "it_name", "es_name", "en_name", "qty"]


@dataclass(frozen=True)
class HeaderMatch:
    sheet_name: str
    header_start_row: int
    header_row: int
    columns: dict[str, int]
    resolution_methods: dict[str, str]
    score: int


class LegacyXlsCell:
    """Small 1-based cell adapter so BIFF workbooks share the XLSX parser."""

    def __init__(self, value: Any, coordinate: str) -> None:
        self.value = value
        self.coordinate = coordinate


class LegacyXlsWorksheet:
    def __init__(self, worksheet: Any) -> None:
        self._worksheet = worksheet
        self.title = worksheet.name
        self.max_row = worksheet.nrows
        self.max_column = worksheet.ncols

    def cell(self, row: int, column: int) -> LegacyXlsCell:
        if row < 1 or column < 1 or row > self.max_row or column > self.max_column:
            value = ""
        else:
            value = self._worksheet.cell_value(row - 1, column - 1)
        return LegacyXlsCell(value, f"{get_column_letter(column)}{row}")


class LegacyXlsWorkbook:
    def __init__(self, workbook: Any) -> None:
        self._workbook = workbook
        self.worksheets = [LegacyXlsWorksheet(sheet) for sheet in workbook.sheets()]
        self._worksheets_by_name = {worksheet.title: worksheet for worksheet in self.worksheets}

    def __getitem__(self, sheet_name: str) -> LegacyXlsWorksheet:
        return self._worksheets_by_name[sheet_name]


@dataclass(frozen=True)
class ImageAsset:
    data: bytes
    extension: str
    source: str
    dispimg_id: str | None = None
    anchor_row: int | None = None
    anchor_col: int | None = None


def contract_media_type(path_or_extension: Path | str) -> str:
    value = str(path_or_extension).lower()
    suffix = value if value in CONTRACT_MEDIA_TYPES else Path(value).suffix
    return CONTRACT_MEDIA_TYPES.get(suffix, "application/octet-stream")


def create_project_id(contract_path: Path) -> str:
    digest = hashlib.sha1(str(contract_path).encode("utf-8")).hexdigest()[:8]
    stem = re.sub(r"[^A-Za-z0-9_-]+", "-", contract_path.stem).strip("-") or "contract"
    # Include microseconds so two uploads of the same contract in one second do
    # not silently target the same project directory.
    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
    return f"{timestamp}-{stem[:36]}-{digest}"


def validate_new_project_id(workspace_root: Path, project_id: str) -> Path:
    if not project_id or project_id in {".", ".."} or not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_.-]{0,79}", project_id):
        raise ValueError(
            "Invalid project_id; use 1-80 letters, digits, dots, underscores or hyphens, starting with a letter or digit."
        )
    projects_root = (workspace_root / "projects").resolve()
    project_dir = (projects_root / project_id).resolve()
    try:
        project_dir.relative_to(projects_root)
    except ValueError as exc:
        raise ValueError("Invalid project_id path.") from exc
    if project_dir.exists():
        raise ValueError(
            f"Project already exists: {project_id}. Load the existing project or choose a new project_id."
        )
    return project_dir


def extract_contract_project(
    contract_path: Path | str,
    workspace_root: Path | str | None = None,
    project_id: str | None = None,
    original_filename: str | None = None,
) -> dict[str, Any]:
    contract_path = Path(contract_path)
    if not contract_path.exists():
        raise FileNotFoundError(f"Contract file not found: {contract_path}")
    extension = contract_path.suffix.lower()
    if extension not in SUPPORTED_CONTRACT_EXTENSIONS:
        allowed = ", ".join(sorted(SUPPORTED_CONTRACT_EXTENSIONS))
        raise ValueError(f"Unsupported contract format: {contract_path.suffix or '(no extension)'}. Supported: {allowed}.")
    if extension == ".xlsx":
        validate_xlsx_container(contract_path)

    source_filename = normalize_original_contract_filename(
        original_filename or contract_path.name,
        extension=extension,
    )
    workspace_root = Path(workspace_root) if workspace_root else workspace_root_from_app()
    project_id = project_id or create_project_id(Path(source_filename))
    project_dir = validate_new_project_id(workspace_root, project_id)
    project_dir.parent.mkdir(parents=True, exist_ok=True)
    try:
        # mkdir(exist_ok=False) is the atomic reservation. The earlier exists
        # check only improves the error message and cannot prevent a race.
        project_dir.mkdir()
    except FileExistsError as exc:
        raise ValueError(
            f"Project already exists: {project_id}. Load the existing project or choose a new project_id."
        ) from exc
    try:
        return _extract_contract_project_into(
            contract_path,
            extension=extension,
            project_id=project_id,
            project_dir=project_dir,
            original_filename=source_filename,
        )
    except Exception:
        # A failed parse must not leave a half-created project that looks
        # loadable or can contaminate a later retry with stale artifacts.
        shutil.rmtree(project_dir, ignore_errors=True)
        raise


def _extract_contract_project_into(
    contract_path: Path,
    *,
    extension: str,
    project_id: str,
    project_dir: Path,
    original_filename: str,
) -> dict[str, Any]:
    input_dir = project_dir / "input"
    data_dir = project_dir / "data"
    reports_dir = project_dir / "reports"
    product_images_dir = project_dir / "extracted" / "product_images"
    for directory in [
        input_dir,
        data_dir,
        reports_dir,
        product_images_dir,
    ]:
        directory.mkdir(parents=True, exist_ok=True)

    copied_contract = input_dir / f"contract{extension}"
    if contract_path.resolve() != copied_contract.resolve():
        shutil.copy2(contract_path, copied_contract)

    result = parse_contract(copied_contract, project_id, project_dir)
    result["report"]["original_filename"] = original_filename
    normalized_path = data_dir / "items.normalized.json"
    contract_snapshot_path = data_dir / "contract.snapshot.json"
    report_path = reports_dir / "extraction_report.json"
    unmatched_path = reports_dir / "unmatched_images.json"
    lineart_groups_path = reports_dir / "lineart_groups.json"

    write_standard_json(normalized_path, result["normalized"], "items.normalized.schema.json")
    write_json(report_path, result["report"])
    write_json(unmatched_path, result["unmatched_images"])
    contract_snapshot = build_contract_snapshot(
        project_id=project_id,
        contract_path=copied_contract,
        extraction_report=result["report"],
        items=result["contract_snapshot_items"],
    )
    write_standard_json(
        contract_snapshot_path,
        contract_snapshot,
        "contract.snapshot.schema.json",
    )
    lineart_groups = build_lineart_groups_report(result["normalized"])
    write_standard_json(
        lineart_groups_path,
        lineart_groups,
        "lineart_groups.schema.json",
    )
    # Contract ingestion stops at canonical fields plus immutable source images.
    # Product-region and line-art workspaces are initialized lazily by their
    # explicit downstream actions.

    result["paths"] = {
        "project_dir": str(project_dir),
        "contract": str(copied_contract),
        "normalized": str(normalized_path),
        "contract_snapshot": str(contract_snapshot_path),
        "report": str(report_path),
        "unmatched_images": str(unmatched_path),
        "lineart_groups": str(lineart_groups_path),
        "lineart_report": str(reports_dir / "lineart_report.json"),
        "product_regions": str(reports_dir / "product_regions.json"),
    }
    return result


def load_contract_workbook(contract_path: Path, *, data_only: bool = False) -> Any:
    """Load one workbook view; legacy XLS already exposes cached values."""

    suffix = contract_path.suffix.lower()
    if suffix == ".xlsx":
        return load_workbook(contract_path, data_only=data_only, read_only=False)
    if suffix == ".xls":
        try:
            import xlrd
        except ImportError as exc:  # pragma: no cover - deployment configuration
            raise ValueError("Legacy .xls support requires xlrd==2.0.2.") from exc
        return LegacyXlsWorkbook(
            xlrd.open_workbook(
                str(contract_path),
                on_demand=False,
                formatting_info=True,
            )
        )
    allowed = ", ".join(sorted(SUPPORTED_CONTRACT_EXTENSIONS))
    raise ValueError(f"Unsupported contract format: {contract_path.suffix or '(no extension)'}. Supported: {allowed}.")


def business_fields_use_formulas(workbook: Any, header: HeaderMatch) -> bool:
    """Return whether canonical business columns require a cached-value view.

    Image formulas such as DISPIMG deliberately do not count: they are read
    from the formula workbook and must not force a second 40+ MB workbook load.
    """

    worksheet = workbook[header.sheet_name]
    business_columns = [header.columns[field] for field in REQUIRED_FIELDS]
    for row in range(header.header_row + 1, worksheet.max_row + 1):
        for column in business_columns:
            cell = worksheet.cell(row, column)
            if getattr(cell, "data_type", None) == "f":
                return True
            value = cell.value
            if isinstance(value, str) and value.lstrip().startswith("="):
                return True
    return False


def validate_xlsx_container(contract_path: Path) -> None:
    """Reject malformed or explosively compressed XLSX containers early."""

    try:
        with ZipFile(contract_path) as archive:
            entries = archive.infolist()
    except BadZipFile as exc:
        raise ValueError("Invalid .xlsx file: the upload is not a valid ZIP workbook.") from exc
    validate_archive_entries(entries, label=".xlsx file")


def validate_archive_entries(entries: list[Any], *, label: str) -> None:
    if len(entries) > MAX_XLSX_ARCHIVE_FILES:
        raise ValueError(f"Invalid {label}: archive contains too many entries ({len(entries)}).")
    total_uncompressed = sum(entry.file_size for entry in entries)
    if total_uncompressed > MAX_XLSX_UNCOMPRESSED_BYTES:
        raise ValueError(f"Invalid {label}: expanded content exceeds the 1 GB safety limit.")
    for entry in entries:
        if entry.file_size > MAX_XLSX_ENTRY_BYTES:
            raise ValueError(f"Invalid {label}: expanded entry is too large: {entry.filename}")
        if entry.file_size < 10 * 1024 * 1024:
            continue
        ratio = entry.file_size / max(entry.compress_size, 1)
        if ratio > MAX_XLSX_COMPRESSION_RATIO:
            raise ValueError(f"Invalid {label}: suspicious compression ratio in {entry.filename}.")


def parse_contract(contract_path: Path, project_id: str, project_dir: Path) -> dict[str, Any]:
    # Keep the formula workbook for DISPIMG and drawing anchors, but read
    # business fields from Excel's cached values.  This matters for contracts
    # that use XLOOKUP/VLOOKUP: parsing the formula text itself produces fake
    # barcodes and missing quantities even though Excel displays valid values.
    workbook = load_contract_workbook(contract_path, data_only=False)
    header = find_best_header(workbook)
    cached_workbook = (
        load_contract_workbook(contract_path, data_only=True)
        if contract_path.suffix.lower() == ".xlsx" and business_fields_use_formulas(workbook, header)
        else workbook
    )
    worksheet = workbook[header.sheet_name]
    cached_worksheet = cached_workbook[header.sheet_name]
    cell_image_assets = load_wps_cell_images(contract_path)
    standard_images_by_row = load_standard_images_by_row(worksheet)
    product_images_dir = project_dir / "extracted" / "product_images"

    warnings: list[dict[str, Any]] = []
    unmatched_images: list[dict[str, Any]] = []
    items: list[dict[str, Any]] = []
    contract_snapshot_items: list[dict[str, Any]] = []
    item_rows: set[int] = set()
    used_dispimg_ids: set[str] = set()
    used_standard_assets: set[str] = set()
    column_letters = {field: get_column_letter(col) for field, col in header.columns.items()}

    row_evidence = []
    for row in range(header.header_row + 1, worksheet.max_row + 1):
        raw_item_no = cell_value(cached_worksheet, worksheet, row, header.columns["item_no"])
        item_no = clean_identifier(raw_item_no)
        raw_barcode = cell_value(cached_worksheet, worksheet, row, header.columns["barcode"])
        qty = normalize_qty(cell_value(cached_worksheet, worksheet, row, header.columns["qty"]))
        names = {
            "cn_name": clean_text(cell_value(cached_worksheet, worksheet, row, header.columns["cn_name"])),
            "it_name": clean_text(cell_value(cached_worksheet, worksheet, row, header.columns["it_name"])),
            "es_name": clean_text(cell_value(cached_worksheet, worksheet, row, header.columns["es_name"])),
            "en_name": clean_text(cell_value(cached_worksheet, worksheet, row, header.columns["en_name"])),
        }
        row_evidence.append(
            inspect_contract_row(
                row=row,
                item_no=item_no,
                barcode_value=raw_barcode,
                qty=qty,
                names=names,
            )
        )

    contract_item_region = identify_contract_item_region(
        row_evidence,
        sheet=header.sheet_name,
        header_rows=[header.header_start_row, header.header_row],
    )

    for evidence in contract_item_region.product_rows:
        row = evidence.row
        item_no = evidence.item_no
        barcode = evidence.barcode.digits
        qty = evidence.qty
        names = dict(evidence.names)
        item_rows.add(row)
        item = {
            "item_no": item_no,
            "barcode": barcode,
            **names,
            "qty": qty,
            "source_image": None,
            "source_image_status": "missing",
            "product_region_status": "missing",
            "selected_product_region": None,
            "lineart_status": "missing",
            "lineart_asset": None,
        }

        validate_item_fields(item, row, warnings)
        image_info = match_row_image(
            worksheet=worksheet,
            row=row,
            header=header,
            item_no=item_no,
            cell_image_assets=cell_image_assets,
            standard_images_by_row=standard_images_by_row,
            product_images_dir=product_images_dir,
            project_dir=project_dir,
            unmatched_images=unmatched_images,
            warnings=warnings,
            used_dispimg_ids=used_dispimg_ids,
            used_standard_assets=used_standard_assets,
        )
        item.update(image_info)
        items.append(item)
        contract_snapshot_items.append(
            snapshot_item(item, source_row=row, columns=column_letters)
        )

    add_duplicate_warnings(items, "item_no", "duplicate_item_no", warnings)
    add_duplicate_warnings(items, "barcode", "duplicate_barcode", warnings)

    if not used_dispimg_ids:
        for row, assets in standard_images_by_row.items():
            if row not in item_rows:
                for asset in assets:
                    unmatched_images.append(
                        {
                            "reason": "standard_image_anchor_row_has_no_item",
                            "source": asset.source,
                            "anchor_row": asset.anchor_row,
                            "anchor_col": asset.anchor_col,
                        }
                    )
                    add_warning(
                        warnings,
                        "unmatched_image",
                        "Standard Excel image could not be matched to an item row.",
                        row=asset.anchor_row,
                    )

    normalized = {"project_id": project_id, "items": items}
    report = {
        "project_id": project_id,
        "source_contract": f"input/contract{contract_path.suffix.lower()}",
        "contract_format": contract_path.suffix.lower().lstrip("."),
        "selected_sheet": header.sheet_name,
        "header_start_row": header.header_start_row,
        "header_row": header.header_row,
        "header_rows": [header.header_start_row, header.header_row],
        "columns": column_letters,
        "column_resolution": dict(header.resolution_methods),
        "contract_item_region": contract_item_region.report(
            sheet=header.sheet_name,
            header_rows=[header.header_start_row, header.header_row],
        ),
        "total_items": len(items),
        "matched_images": sum(1 for item in items if item["source_image_status"] == "matched"),
        "missing_images": [item["item_no"] for item in items if item["source_image_status"] == "missing"],
        "unmatched_images_count": len(unmatched_images),
        "unused_cell_images_count": max(len(cell_image_assets) - len(used_dispimg_ids), 0),
        "warnings": warnings,
        "errors": [],
    }
    return {
        "normalized": normalized,
        "contract_snapshot_items": contract_snapshot_items,
        "report": report,
        "unmatched_images": {"project_id": project_id, "images": unmatched_images},
    }


def cell_value(cached_worksheet: Any, formula_worksheet: Any, row: int, column: int) -> Any:
    """Prefer Excel's cached result and fall back to the formula cell value."""

    cached = cached_worksheet.cell(row, column).value
    return cached if cached is not None else formula_worksheet.cell(row, column).value


def find_best_header(workbook: Any) -> HeaderMatch:
    candidates: list[HeaderMatch] = []
    partial_candidates: list[HeaderMatch] = []
    for worksheet in workbook.worksheets:
        scan_rows = min(worksheet.max_row, 30)
        for header_row in range(1, scan_rows + 1):
            for span in range(1, min(3, header_row) + 1):
                header_start_row = header_row - span + 1
                alias_columns = detect_columns(
                    worksheet,
                    header_start_row,
                    header_row,
                )
                columns = dict(alias_columns)
                # Content profiling is a fallback for a plausible header row,
                # not a workbook-wide search. Requiring explicit aliases first
                # keeps large contracts fast and prevents data rows becoming
                # headers. Data always begins after the final header row.
                content_columns = (
                    repair_safe_columns_by_content(worksheet, header_row, columns)
                    if len(alias_columns) >= 3
                    else {}
                )
                columns.update(content_columns)
                alias_method = "header_alias" if span == 1 else "header_alias_multiline"
                resolution_methods = {
                    **{field: alias_method for field in alias_columns},
                    **{field: "content_profile" for field in content_columns},
                }
                found_required = sum(1 for field in REQUIRED_FIELDS if field in columns)
                image_col = detect_image_column(
                    worksheet,
                    header_start_row,
                    header_row,
                    columns,
                )
                if image_col:
                    columns["image"] = image_col
                    resolution_methods["image"] = "header_alias_or_image_formula"
                # Prefer the smallest span when two candidates resolve the same
                # fields so a one-row header never consumes the first data row.
                score = (
                    len(alias_columns) * 12
                    + len(content_columns) * 7
                    - header_row
                    - (span - 1) * 3
                )
                if image_col:
                    score += 5
                match = HeaderMatch(
                    worksheet.title,
                    header_start_row,
                    header_row,
                    columns,
                    resolution_methods,
                    score,
                )
                if len(alias_columns) >= 2:
                    partial_candidates.append(match)
                if found_required >= 5 and "item_no" in columns:
                    candidates.append(match)
    if not candidates:
        best_partial = max(partial_candidates, key=lambda candidate: candidate.score, default=None)
        raise ValueError(header_mapping_error(workbook, best_partial))
    candidates.sort(key=lambda candidate: candidate.score, reverse=True)
    best = candidates[0]
    missing = [field for field in REQUIRED_FIELDS if field not in best.columns]
    if missing:
        raise ValueError(header_mapping_error(workbook, best))
    return best


def detect_columns(
    worksheet: Any,
    header_start_row: int,
    header_row: int | None = None,
) -> dict[str, int]:
    header_row = header_start_row if header_row is None else header_row
    columns: dict[str, int] = {}
    for col in range(1, worksheet.max_column + 1):
        field = canonical_field_for_header_parts(
            header_label_parts(worksheet, header_start_row, header_row, col)
        )
        if field and field not in columns:
            columns[field] = col
    return columns


def canonical_field_for_header_parts(parts: list[str]) -> str | None:
    normalized_parts = [normalize_label(part) for part in parts if normalize_label(part)]
    if not normalized_parts:
        return None
    candidates = [
        "".join(normalized_parts),
        "".join(reversed(normalized_parts)),
        *reversed(normalized_parts),
    ]
    seen: set[str] = set()
    for candidate in candidates:
        if not candidate or candidate in seen:
            continue
        seen.add(candidate)
        field = canonical_field_for_header(candidate)
        if field:
            return field
    return None


def header_label_parts(
    worksheet: Any,
    header_start_row: int,
    header_row: int,
    col: int,
) -> list[str]:
    parts: list[str] = []
    for row in range(header_start_row, header_row + 1):
        text = clean_text(
            header_cell_value(
                worksheet,
                row,
                col,
                minimum_anchor_row=header_start_row,
            )
        )
        if text and (not parts or normalize_label(text) != normalize_label(parts[-1])):
            parts.append(text)
    return parts


def header_cell_value(
    worksheet: Any,
    row: int,
    col: int,
    *,
    minimum_anchor_row: int,
) -> Any:
    value = worksheet.cell(row, col).value
    if value is not None and clean_text(value):
        return value
    for min_row, max_row, min_col, max_col in merged_cell_bounds(worksheet):
        if (
            min_row >= minimum_anchor_row
            and min_row <= row <= max_row
            and min_col <= col <= max_col
        ):
            return worksheet.cell(min_row, min_col).value
    return value


def merged_cell_bounds(worksheet: Any) -> list[tuple[int, int, int, int]]:
    bounds: list[tuple[int, int, int, int]] = []
    merged_cells = getattr(worksheet, "merged_cells", None)
    for merged_range in getattr(merged_cells, "ranges", ()):
        bounds.append(
            (
                int(merged_range.min_row),
                int(merged_range.max_row),
                int(merged_range.min_col),
                int(merged_range.max_col),
            )
        )
    legacy_sheet = getattr(worksheet, "_worksheet", None)
    for row_low, row_high, col_low, col_high in getattr(legacy_sheet, "merged_cells", ()):
        bounds.append((row_low + 1, row_high, col_low + 1, col_high))
    return bounds


def canonical_field_for_header(normalized_label: str) -> str | None:
    """Resolve explicit/decorated aliases using longest-match semantics.

    Longest match is essential because the generic Chinese alias ``品名`` is
    contained in language-specific labels such as ``意文品名``.
    """

    exact: dict[str, set[str]] = {}
    matches: list[tuple[int, str]] = []
    for field, aliases in FIELD_ALIASES.items():
        for alias in aliases:
            normalized_alias = normalize_label(alias)
            if not normalized_alias:
                continue
            exact.setdefault(normalized_alias, set()).add(field)
            if normalized_alias in normalized_label:
                matches.append((len(normalized_alias), field))
    exact_fields = exact.get(normalized_label, set())
    if len(exact_fields) == 1:
        return next(iter(exact_fields))
    if not matches:
        return None
    longest = max(length for length, _ in matches)
    fields = {field for length, field in matches if length == longest}
    return next(iter(fields)) if len(fields) == 1 else None


def repair_safe_columns_by_content(
    worksheet: Any,
    header_row: int,
    columns: dict[str, int],
) -> dict[str, int]:
    """Recover only fields whose values have a defensible structural profile.

    Foreign-language name columns are intentionally excluded: similar Latin
    text cannot prove whether a column is Italian, Spanish, or English.
    """

    repaired: dict[str, int] = {}
    used = set(columns.values())
    for field in ("barcode", "item_no", "qty", "cn_name"):
        if field in columns:
            continue
        scored = []
        for col in range(1, worksheet.max_column + 1):
            if col in used:
                continue
            score = content_profile_score(worksheet, header_row, col, field)
            scored.append((score, col))
        scored.sort(reverse=True)
        if not scored:
            continue
        best_score, best_col = scored[0]
        second_score = scored[1][0] if len(scored) > 1 else 0.0
        if best_score >= 0.75 and best_score - second_score >= 0.20:
            repaired[field] = best_col
            used.add(best_col)
    return repaired


def content_profile_score(worksheet: Any, header_row: int, col: int, field: str) -> float:
    values = []
    end_row = min(worksheet.max_row, header_row + 40)
    for row in range(header_row + 1, end_row + 1):
        value = worksheet.cell(row, col).value
        if value is not None and clean_text(value):
            values.append(value)
        if len(values) >= 20:
            break
    if not values:
        return 0.0

    def matches(value: Any) -> bool:
        text = clean_text(value)
        compact = clean_identifier(value)
        digits = re.sub(r"\D", "", text)
        if field == "barcode":
            return len(digits) == 13 and (is_valid_ean13(digits) or digits == "0" * 13)
        if field == "item_no":
            if digits == compact and len(digits) in {12, 13, 14}:
                return False
            return bool(re.fullmatch(r"[A-Za-z]{0,8}[-_]?\d{3,10}(?:[-_][A-Za-z0-9]+)?", compact))
        if field == "qty":
            quantity = normalize_qty(value)
            return quantity is not None and 0 < float(quantity) <= 10_000 and len(digits) <= 5
        if field == "cn_name":
            return bool(re.search(r"[\u3400-\u4dbf\u4e00-\u9fff]", text))
        return False

    return sum(1 for value in values if matches(value)) / len(values)


def header_mapping_error(workbook: Any, match: HeaderMatch | None) -> str:
    if match is None:
        return (
            "header_mapping_required: no defensible contract header candidate was found; "
            f"required={','.join(REQUIRED_FIELDS)}"
        )
    worksheet = workbook[match.sheet_name]
    missing = [field for field in REQUIRED_FIELDS if field not in match.columns]
    used = set(match.columns.values())
    unresolved = []
    for col in range(1, worksheet.max_column + 1):
        raw = " / ".join(
            header_label_parts(
                worksheet,
                match.header_start_row,
                match.header_row,
                col,
            )
        )
        if raw and col not in used:
            unresolved.append(f"{get_column_letter(col)}:{raw!r}")
    recognized = ",".join(
        f"{field}={get_column_letter(col)}" for field, col in sorted(match.columns.items())
    )
    return (
        "header_mapping_required: "
        f"missing={','.join(missing)}; "
        f"sheet={match.sheet_name!r}; rows={match.header_start_row}-{match.header_row}; "
        f"recognized={recognized or '(none)'}; "
        f"candidates={','.join(unresolved) or '(none)'}"
    )


def detect_image_column(
    worksheet: Any,
    header_start_row: int,
    header_row: int,
    columns: dict[str, int],
) -> int | None:
    image_aliases = {normalize_label(alias) for alias in IMAGE_ALIASES}
    for col in range(1, worksheet.max_column + 1):
        labels = [
            normalize_label(part)
            for part in header_label_parts(worksheet, header_start_row, header_row, col)
        ]
        labels.extend(["".join(labels), "".join(reversed(labels))])
        if any(
            label and any(alias and alias in label for alias in image_aliases)
            for label in labels
        ):
            return col

    item_col = columns.get("item_no")
    barcode_col = columns.get("barcode")
    if item_col and barcode_col and abs(barcode_col - item_col) > 1:
        left, right = sorted([item_col, barcode_col])
        for col in range(left + 1, right):
            if column_has_dispimg(worksheet, header_row + 1, col):
                return col

    field_cols = set(columns.values())
    for col in range(1, worksheet.max_column + 1):
        if col in field_cols:
            continue
        if column_has_dispimg(worksheet, header_row + 1, col):
            return col
    return None


def column_has_dispimg(worksheet: Any, start_row: int, col: int) -> bool:
    end_row = min(worksheet.max_row, start_row + 20)
    for row in range(start_row, end_row + 1):
        if extract_dispimg_id(worksheet.cell(row, col).value):
            return True
    return False


def match_row_image(
    worksheet: Any,
    row: int,
    header: HeaderMatch,
    item_no: str,
    cell_image_assets: dict[str, ImageAsset],
    standard_images_by_row: dict[int, list[ImageAsset]],
    product_images_dir: Path,
    project_dir: Path,
    unmatched_images: list[dict[str, Any]],
    warnings: list[dict[str, Any]],
    used_dispimg_ids: set[str],
    used_standard_assets: set[str],
) -> dict[str, Any]:
    dispimg_cell = None
    dispimg_id = None
    image_col = header.columns.get("image")
    if image_col:
        dispimg_cell = worksheet.cell(row, image_col)
        dispimg_id = extract_dispimg_id(dispimg_cell.value)

    if not dispimg_id:
        for col in range(1, worksheet.max_column + 1):
            cell = worksheet.cell(row, col)
            candidate_id = extract_dispimg_id(cell.value)
            if candidate_id:
                dispimg_cell = cell
                dispimg_id = candidate_id
                break

    if dispimg_id:
        asset = cell_image_assets.get(dispimg_id)
        if asset:
            used_dispimg_ids.add(dispimg_id)
            return save_matched_image(asset, item_no, product_images_dir, project_dir)
        cell_ref = dispimg_cell.coordinate if dispimg_cell is not None else None
        unmatched_images.append(
            {
                "reason": "dispimg_id_not_found_in_cellimages",
                "item_no": item_no,
                "row": row,
                "cell": cell_ref,
                "dispimg_id": dispimg_id,
            }
        )
        add_warning(
            warnings,
            "unmatched_image",
            f"Image formula {dispimg_id} could not be resolved from workbook cell images.",
            item_no=item_no,
            row=row,
            cell=cell_ref,
        )
        return {"source_image": None, "source_image_status": "unmatched"}

    standard_assets = standard_images_by_row.get(row, [])
    if standard_assets:
        asset = standard_assets[0]
        used_standard_assets.add(asset.source)
        return save_matched_image(asset, item_no, product_images_dir, project_dir)

    add_warning(
        warnings,
        "missing_product_image",
        "Item has no matched product image.",
        item_no=item_no,
        row=row,
    )
    return {"source_image": None, "source_image_status": "missing"}


def save_matched_image(
    asset: ImageAsset,
    item_no: str,
    product_images_dir: Path,
    project_dir: Path,
) -> dict[str, Any]:
    filename = unique_image_filename(product_images_dir, item_no, asset.extension)
    output_path = product_images_dir / filename
    output_path.write_bytes(asset.data)
    return {
        "source_image": output_path.relative_to(project_dir).as_posix(),
        "source_image_status": "matched",
    }


def unique_image_filename(directory: Path, item_no: str, extension: str) -> str:
    safe_item = re.sub(r"[^A-Za-z0-9_.-]+", "_", item_no).strip("._") or "item"
    extension = extension if extension.startswith(".") else f".{extension}"
    candidate = f"{safe_item}{extension.lower()}"
    counter = 2
    while (directory / candidate).exists():
        candidate = f"{safe_item}_{counter}{extension.lower()}"
        counter += 1
    return candidate


def load_wps_cell_images(contract_path: Path) -> dict[str, ImageAsset]:
    archive_bytes = read_wps_cell_image_archive(contract_path)
    if not archive_bytes:
        return {}

    try:
        archive = ZipFile(BytesIO(archive_bytes))
    except BadZipFile:
        return {}

    with archive:
        validate_archive_entries(archive.infolist(), label="embedded WPS image package")
        names = set(archive.namelist())
        cell_images_name = find_archive_name(names, "xl/cellimages.xml")
        relationships_name = find_archive_name(names, "xl/_rels/cellimages.xml.rels")
        if not cell_images_name or not relationships_name:
            return {}

        rel_root = ET.fromstring(archive.read(relationships_name))
        rels = {rel.attrib.get("Id"): rel.attrib.get("Target") for rel in rel_root}
        root = ET.fromstring(archive.read(cell_images_name))
        assets: dict[str, ImageAsset] = {}
        for pic in root.findall(".//xdr:pic", NAMESPACES):
            c_nv_pr = pic.find(".//xdr:cNvPr", NAMESPACES)
            blip = pic.find(".//a:blip", NAMESPACES)
            if c_nv_pr is None or blip is None:
                continue
            dispimg_id = c_nv_pr.attrib.get("name")
            rel_id = blip.attrib.get(f"{{{NAMESPACES['r']}}}embed")
            target = rels.get(rel_id or "")
            if not dispimg_id or not target or target.upper() == "NULL":
                continue
            archive_path = posixpath.normpath(posixpath.join("xl", target))
            actual_archive_path = find_archive_name(names, archive_path)
            if not actual_archive_path:
                continue
            extension = Path(actual_archive_path).suffix or ".png"
            assets[dispimg_id] = ImageAsset(
                data=archive.read(actual_archive_path),
                extension=extension,
                source=actual_archive_path,
                dispimg_id=dispimg_id,
            )
        return assets


def read_wps_cell_image_archive(contract_path: Path) -> bytes | None:
    """Read the WPS cell-image package from XLSX or an OLE-based XLS file."""

    if contract_path.suffix.lower() == ".xlsx":
        return contract_path.read_bytes()
    if contract_path.suffix.lower() != ".xls":
        return None

    try:
        import olefile
    except ImportError as exc:  # pragma: no cover - deployment configuration
        raise ValueError("Legacy .xls image extraction requires olefile==0.47.") from exc

    ole = olefile.OleFileIO(str(contract_path))
    try:
        if not ole.exists("ETCellImageData"):
            return None
        return ole.openstream("ETCellImageData").read()
    finally:
        ole.close()


def find_archive_name(names: set[str], expected: str) -> str | None:
    expected = expected.replace("\\", "/").lower()
    return next((name for name in names if name.lower() == expected), None)


def load_standard_images_by_row(worksheet: Any) -> dict[int, list[ImageAsset]]:
    images_by_row: dict[int, list[ImageAsset]] = {}
    for index, image in enumerate(getattr(worksheet, "_images", [])):
        try:
            marker = image.anchor._from
            row = marker.row + 1
            col = marker.col + 1
            extension = f".{getattr(image, 'format', 'png') or 'png'}"
            asset = ImageAsset(
                data=image._data(),
                extension=extension,
                source=f"{worksheet.title}:image{index + 1}",
                anchor_row=row,
                anchor_col=col,
            )
            images_by_row.setdefault(row, []).append(asset)
        except Exception:
            continue
    return images_by_row


def validate_item_fields(item: dict[str, Any], row: int, warnings: list[dict[str, Any]]) -> None:
    if not item["item_no"]:
        add_warning(warnings, "missing_item_no", "Item number is required.", row=row)
    if not re.fullmatch(r"\d{13}", item["barcode"] or ""):
        add_warning(
            warnings,
            "invalid_barcode_format",
            "Barcode must contain exactly 13 digits after normalization.",
            item_no=item["item_no"],
            row=row,
            value=item["barcode"],
        )
    elif not is_valid_ean13(item["barcode"]):
        add_warning(
            warnings,
            "invalid_ean13_checksum",
            "Barcode failed EAN-13 checksum validation.",
            item_no=item["item_no"],
            row=row,
            value=item["barcode"],
        )
    for field in ["cn_name", "it_name", "es_name", "en_name"]:
        if not item[field]:
            add_warning(
                warnings,
                f"missing_{field}",
                f"{field} is required.",
                item_no=item["item_no"],
                row=row,
            )
    if item["qty"] is None:
        add_warning(
            warnings,
            "invalid_qty",
            "Quantity must be numeric.",
            item_no=item["item_no"],
            row=row,
        )


def add_duplicate_warnings(
    items: list[dict[str, Any]],
    field: str,
    code: str,
    warnings: list[dict[str, Any]],
) -> None:
    values = [item[field] for item in items if item.get(field)]
    for value, count in Counter(values).items():
        if count > 1:
            add_warning(
                warnings,
                code,
                f"{field} appears {count} times.",
                value=value,
            )


def add_warning(
    warnings: list[dict[str, Any]],
    code: str,
    message: str,
    **details: Any,
) -> None:
    warning = {"severity": "warning", "code": code, "message": message}
    warning.update({key: value for key, value in details.items() if value is not None})
    warnings.append(warning)


def extract_dispimg_id(value: Any) -> str | None:
    if not isinstance(value, str):
        return None
    match = DISPIMG_RE.search(value)
    return match.group(1) if match else None


def normalize_label(value: Any) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value)).casefold()
    # Keep letters/numbers from every language and remove layout punctuation,
    # units' brackets, slashes, underscores and whitespace uniformly.
    return re.sub(r"[\W_]+", "", text, flags=re.UNICODE)


def clean_identifier(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", "", str(value).strip())


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).strip())


def normalize_qty(value: Any) -> int | float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value) if value.is_integer() else value
    text = str(value).strip().replace(",", "")
    if re.fullmatch(r"\d+(?:\.\d+)?(?:\s*[*xX×]\s*\d+(?:\.\d+)?)+", text):
        product = 1.0
        for part in re.split(r"[*xX×]", text):
            product *= float(part.strip())
        return int(product) if product.is_integer() else product
    try:
        number = float(text)
    except ValueError:
        return None
    return int(number) if number.is_integer() else number


def write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
